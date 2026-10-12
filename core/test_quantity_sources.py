import json
import hashlib
import importlib
import importlib.util
import io
import shutil
import tempfile
import zipfile
from types import SimpleNamespace
from unittest.mock import patch
import ezdxf
from copy import deepcopy
from pathlib import Path
from unittest import skipUnless
from django.test import SimpleTestCase, RequestFactory, override_settings
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from core import quantity_views as views


def small_pdf():
    writer = PdfWriter()
    for title, number in [('COLUMN SCHEDULE', 'ST-101'), ('FRAMING PLAN', 'ST-102')]:
        page = writer.add_blank_page(width=600, height=800)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                 NameObject('/Subtype'): NameObject('/Type1'),
                                 NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({
            NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(('BT /F1 12 Tf 40 80 Td (DRAWING TITLE: ' + title +
                         ') Tj 0 -20 Td (DRAWING NO: ' + number + ') Tj ET').encode('ascii'))
        page[NameObject('/Contents')] = writer._add_object(stream)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


class QuantitySourceManifestTests(SimpleTestCase):
    def setUp(self):
        self.assertIsNotNone(importlib.util.find_spec('core.quantity_sources'),
                             'source identity module must exist')
        self.sources = importlib.import_module('core.quantity_sources')

    def manifest(self, pages, cad=None, total=None):
        return self.sources.build_source_manifest(
            pages, document_sha='pdf-sha', total_pages=total or len(pages),
            cad_data=cad or {})

    def test_explicit_spaced_title_identifies_schedule(self):
        m = self.manifest({1: '도 면 명 : 기 둥 일 람 표\n도 면 번 호 : ST-001'})
        row = m['pdf_pages'][0]
        self.assertEqual(row['role'], 'column_schedule')
        self.assertEqual(row['drawing_number'], 'ST-001')
        self.assertEqual(row['identity_status'], 'identified')
        self.assertEqual(m['context_pages'], [1])
        self.assertEqual(row['source_id'], 'pdf:pdf-sha:1')

    def test_body_reference_does_not_become_sheet_identity(self):
        m = self.manifest({1: '보는 ST-501 보일람표 참조\nA1 D16'})
        self.assertIsNone(m['pdf_pages'][0]['drawing_number'])
        self.assertEqual(m['context_pages'], [])
        self.assertEqual(m['pdf_pages'][0]['identity_status'], 'unidentified')

    def test_index_is_not_a_physical_schedule_context(self):
        m = self.manifest({1: '도면명: 도면목록\n도면번호: ST-000\nST-101 기둥일람표\nST-102 구조평면도'})
        self.assertEqual(m['pdf_pages'][0]['role'], 'drawing_index')
        self.assertEqual(m['context_pages'], [])

    def test_missing_text_pages_are_visible(self):
        m = self.manifest({1: '도면명: 구조평면도\n도면번호: ZX-034'}, total=3)
        self.assertEqual(len(m['pdf_pages']), 3)
        self.assertEqual([r['page'] for r in m['pdf_pages'] if r['identity_status'] == 'unidentified'], [2, 3])
        self.assertFalse(m['identity_complete'])

    def test_conflicting_explicit_numbers_do_not_match(self):
        m = self.manifest({1: '도면명: 보일람표\n도면번호: ST-100\n도면번호: ST-200'},
                          {'x.dxf': {'texts': ['도면명: 보일람표', '도면번호: ST-100'], 'source_member_sha': 'cad'}})
        self.assertEqual(m['pdf_pages'][0]['identity_status'], 'conflict')
        self.assertEqual(m['links'], [])

    def test_unique_identity_links_without_filename_inference(self):
        m = self.manifest({1: 'DRAWING TITLE: FRAMING PLAN\nDRAWING NO: ZX-034'},
                          {'unrelated-name.dxf': {'texts': ['DRAWING TITLE: FRAMING PLAN', 'DRAWING NO: ZX-034'],
                                                  'source_member_sha': 'cad-sha'}})
        self.assertEqual(len(m['links']), 1)
        self.assertEqual(m['links'][0]['cad_file'], 'unrelated-name.dxf')
        self.assertEqual(m['links'][0]['pdf_page'], 1)
        self.assertEqual(m['links'][0]['cad_source_id'], 'cad:cad-sha:unrelated-name.dxf')

    def test_duplicate_cad_id_remains_unresolved(self):
        cad = {n: {'texts': ['도면명: 구조평면도', '도면번호: ST-100'], 'source_member_sha': n}
               for n in ['a.dxf', 'b.dxf']}
        m = self.manifest({1: '도면명: 구조평면도\n도면번호: ST-100'}, cad)
        self.assertEqual(m['links'], [])
        self.assertTrue(any(i['reason'] == 'ambiguous_pdf_cad_identity' for i in m['issues']))

    def test_same_number_with_different_roles_is_not_matched(self):
        m = self.manifest({1: '도면명: 기둥일람표\n도면번호: ST-100'},
                          {'a.dxf': {'texts': ['도면명: 보일람표', '도면번호: ST-100'], 'source_member_sha': 'cad'}})
        self.assertEqual(m['links'], [])
        self.assertTrue(any(i['reason'] == 'pdf_cad_role_conflict' for i in m['issues']))

    def test_filename_alone_cannot_supply_title_identity(self):
        m = self.manifest({1: '도면명: 구조평면도\n도면번호: ST-100'},
                          {'ST-100_구조평면도.dxf': {'texts': [], 'source_member_sha': 'cad'}})
        self.assertEqual(m['links'], [])
        self.assertEqual(m['cad_files'][0]['identity_status'], 'unidentified')

    def test_sources_locate_candidates_without_claiming_verified_quantities(self):
        m = self.manifest({1: '도면명: 구조평면도\n도면번호: ST-100'})
        original = {'columns': [{'mark': 'C1', 'bbox': {'page': 1, 'x': .1, 'y': .2, 'w': .1, 'h': .1}}]}
        saved = deepcopy(original)
        result = self.sources.attach_member_sources(original, m)
        item = result['columns'][0]
        self.assertEqual(item['source_refs'][0]['page'], 1)
        self.assertEqual(item['source_review']['status'], 'located')
        self.assertFalse(item['source_review']['quantity_verified'])
        self.assertEqual(original, saved)

    def test_schedule_bbox_warns_that_definition_is_not_an_instance(self):
        m = self.manifest({1: '도면명: 기둥일람표\n도면번호: ST-100'})
        result = self.sources.attach_member_sources({'columns': [{'mark': 'C1', 'count': 8, 'bbox': {'page': 1}}]}, m)
        self.assertEqual(result['columns'][0]['source_review']['status'], 'definition_only')
        self.assertEqual(result['columns'][0]['count'], 8)
        self.assertIn('개수', result['columns'][0]['source_review']['note'])

    def test_invalid_or_missing_pages_cannot_create_source_references(self):
        m = self.manifest({1: '도면명: 구조평면도\n도면번호: ST-100'})
        for page in [0, 9, True, '1', 1.5, None]:
            with self.subTest(page=page):
                result = self.sources.attach_member_sources({'beams': [{'bbox': {'page': page}}]}, m)
                self.assertEqual(result['beams'][0]['source_refs'], [])
                self.assertNotEqual(result['beams'][0]['source_review']['status'], 'located')


class QuantitySourceIntegrationTests(SimpleTestCase):
    def member(self, page=1, x=.1, count=2):
        return {'mark': 'C1', 'zone': '1F', 'section': 'A',
                'width_m': .4, 'depth_m': .5, 'height_m': 3, 'count': count,
                'bbox': {'page': page, 'x': x, 'y': .2, 'w': .1, 'h': .1}}

    def test_same_mark_at_different_pages_is_not_dropped(self):
        rows, removed, flags = views._dedupe_category_items([self.member(1), self.member(2)], 'columns')
        self.assertEqual(len(rows), 2)
        self.assertEqual(removed, 0)

    def test_same_mark_at_different_locations_is_not_dropped(self):
        rows, removed, flags = views._dedupe_category_items([self.member(x=.1), self.member(x=.8)], 'columns')
        self.assertEqual(len(rows), 2)
        self.assertEqual(removed, 0)

    def test_same_location_with_conflicting_count_requires_review(self):
        rows, removed, flags = views._dedupe_category_items([self.member(count=2), self.member(count=5)], 'columns')
        self.assertEqual(len(rows), 2)
        self.assertEqual(removed, 0)
        self.assertTrue(all(row.get('_dedup_flag') for row in rows))

    def test_same_location_and_count_is_one_repeated_extraction(self):
        rows, removed, flags = views._dedupe_category_items([self.member(), self.member()], 'columns')
        self.assertEqual(len(rows), 1)
        self.assertEqual(removed, 1)

    @skipUnless(shutil.which('pdftotext'), 'Poppler required')
    def test_actual_pdf_source_manifest_records_pages_and_schedule_context(self):
        self.assertTrue(callable(getattr(views, '_build_source_manifest', None)))
        pdf = small_pdf()
        manifest = views._build_source_manifest({}, pdf)
        self.assertEqual(manifest['context_pages'], [1])
        self.assertEqual([r['drawing_number'] for r in manifest['pdf_pages']], ['ST-101', 'ST-102'])
        self.assertEqual(manifest['document_sha'], hashlib.sha256(pdf).hexdigest())

    def test_checklist_keeps_evidence_warning_and_reference(self):
        member = self.member()
        member['source_refs'] = [{'page': 1, 'source_id': 'pdf:sha:1', 'drawing_number': 'ST-101'}]
        member['source_review'] = {'status': 'definition_only', 'note': '일람표 개수 확인'}
        row = views._build_review_checklist({'columns': [member]})[0]
        self.assertEqual(row.get('source_refs'), member['source_refs'])
        self.assertEqual(row.get('source_review'), member['source_review'])

    def test_excel_contains_saved_drawing_and_member_sources(self):
        from openpyxl import load_workbook
        data = {'source_manifest': {'pdf_pages': [{'page': 1, 'source_id': 'pdf:sha:1',
                 'drawing_number': 'ST-101', 'drawing_title': '=danger', 'role': 'column_schedule',
                 'identity_status': 'identified'}], 'cad_files': [], 'links': [], 'issues': []},
                 '_raw_members': {'columns': [{'mark': 'C1', 'source_refs': [{'page': 1,
                  'source_id': 'pdf:sha:1'}], 'source_review': {'status': 'definition_only', 'note': '개수 확인'}}]}}
        request = RequestFactory().post('/api/quantity/download-excel/',
                                        json.dumps({'results': {'structural': data}}),
                                        content_type='application/json')
        response = views._build_excel_response(request)
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(io.BytesIO(response.content))
        self.assertIn('도면 출처', wb.sheetnames)
        self.assertIn('부재 출처', wb.sheetnames)
        values = list(wb['도면 출처'].values)
        self.assertTrue(any('ST-101' in row for row in values))
        self.assertTrue(all(cell.data_type != 'f' for row in wb['도면 출처'] for cell in row))
        self.assertTrue(any('definition_only' in row for row in wb['부재 출처'].values))


class QuantitySourcePersistenceTests(SimpleTestCase):
    def test_confirmation_and_reopen_keep_member_and_document_evidence(self):
        from core.quantity_sources import build_source_manifest, attach_member_sources
        from core.quantity_state import QuantityStateStore
        with tempfile.TemporaryDirectory() as folder, override_settings(QUANTITY_STATE_DB=str(Path(folder) / 'state.sqlite3')):
            store = QuantityStateStore(Path(folder) / 'state.sqlite3')
            store.begin_job('source-job', '7', 'api_run_quantity', 'input')
            manifest = build_source_manifest({1: '도면명: 기둥일람표\n도면번호: ST-101',
                                               2: '도면명: 구조평면도\n도면번호: ST-102'},
                                              document_sha='sha', total_pages=2)
            members = attach_member_sources({'foundations': [{'mark': 'F1', 'length_m': 2,
                    'width_m': 3, 'height_m': .5, 'count': 1, 'bbox': {'page': 2}}]}, manifest)
            views._extraction_store_set('source-job', members, None)
            req = RequestFactory().post('/api/quantity/confirm-review/',
                        {'job_id': 'source-job', 'corrections': json.dumps([{
                           'review_id': 'foundations:0', 'action': 'edit', 'fields': {'count': 2}}])})
            req.user = SimpleNamespace(pk=7, is_staff=True, is_superuser=False, is_authenticated=True)
            response = views.api_quantity_confirm_review(req)
            self.assertEqual(response.status_code, 200, response.content)
            result = json.loads(response.content)['results']['structural']
            self.assertEqual(result['source_manifest'], manifest)
            item = result['_raw_members']['foundations'][0]
            self.assertEqual(item['source_refs'], members['foundations'][0]['source_refs'])
            self.assertEqual(item['count'], 2)
            saved = QuantityStateStore(store.path).get('result', 'source-job')
            self.assertEqual(saved['results']['structural']['source_manifest'], manifest)
            self.assertEqual(saved['results']['structural']['_raw_members']['foundations'][0]['source_refs'], item['source_refs'])


class QuantityCadSourceTests(SimpleTestCase):
    def test_original_zip_member_hash_and_text_handles_are_retained(self):
        doc = ezdxf.new()
        doc.modelspace().add_text('DRAWING TITLE: FRAMING PLAN', dxfattribs={'insert': (20, 40)})
        doc.modelspace().add_text('DRAWING NO: ZX-034', dxfattribs={'insert': (20, 20)})
        stream = io.StringIO()
        doc.write(stream)
        source = stream.getvalue().encode('utf-8')
        bundle = io.BytesIO()
        with zipfile.ZipFile(bundle, 'w') as archive:
            archive.writestr('folder/x.dxf', source)
        parsed = views.parse_dwg_from_zip(bundle.getvalue())['folder/x.dxf']
        self.assertEqual(parsed.get('source_member_sha'), hashlib.sha256(source).hexdigest())
        self.assertEqual(parsed.get('normalized_member_path'), 'folder/x.dxf')
        self.assertEqual(parsed.get('source_texts', [])[0]['insert'][:2], [20.0, 40.0])
        self.assertTrue(parsed['source_texts'][0]['handle'])

    def test_gemini_does_not_receive_large_raw_source_index(self):
        calls = []
        def generate(**kwargs):
            calls.append(kwargs['contents'])
            return SimpleNamespace(text='{}')
        client = SimpleNamespace(models=SimpleNamespace(generate_content=generate))
        data = {'a.dxf': {'texts': ['C1'], 'source_texts': [
                    {'text': 'RAW_METADATA_ONLY' * 200, 'handle': str(i)} for i in range(5000)]}}
        views._extract_structural_members_one_batch(client, data, [], 1, 1)
        text = calls[0][0]
        self.assertNotIn('RAW_METADATA_ONLY', text)
        self.assertIn('C1', text)


class QuantityLegacyDiagnosticTests(SimpleTestCase):
    def test_audit_reports_missing_classification_and_source_without_changing_input(self):
        self.assertIsNotNone(importlib.util.find_spec('core.quantity_legacy_audit'))
        from core.quantity_legacy_audit import diagnose_workspace
        workspace = {'scene': {'objects': [
           {'id': 'a', 'layer': 'original', 'review_group_id': None},
           {'id': 'b', 'layer': 'shadow', 'review_group_id': 'g'}],
           'review_groups': [{'blocker': 'review_required', 'technical_only': False}]},
           'takeoff_rows': [{'object_id': 'a', 'classification_status': 'unclassified',
                             'source_evidence_status': 'missing', 'quantity_eligibility': False},
                            {'object_id': 'b', 'classification_status': 'unclassified',
                             'source_evidence_status': 'missing', 'quantity_eligibility': False}]}
        before = deepcopy(workspace)
        report = diagnose_workspace(workspace)
        self.assertEqual(report['objects_without_group'], 1)
        self.assertEqual(report['unclassified_objects'], 2)
        self.assertEqual(report['source_evidence_missing'], 2)
        self.assertEqual(report['groups_by_blocker'], {'review_required': 1})
        self.assertFalse(report['ready_for_quantity'])
        self.assertEqual(workspace, before)


class QuantitySourceReviewRegressionTests(SimpleTestCase):
    def manifest(self, text, cad=None):
        from core.quantity_sources import build_source_manifest
        return build_source_manifest({1: text}, document_sha='sha', total_pages=1, cad_data=cad or {})

    def test_notes_are_not_number_labels(self):
        for label in ['DRAWING NOTES', 'SHEET NOTES']:
            with self.subTest(label=label):
                manifest = self.manifest('DRAWING TITLE: FRAMING PLAN\n' + label + ': Refer to ST-101')
                self.assertNotEqual(manifest['pdf_pages'][0]['identity_status'], 'identified')
                self.assertIsNone(manifest['pdf_pages'][0]['drawing_number'])

    def test_two_floor_titles_with_same_role_are_conflicting(self):
        manifest = self.manifest('DRAWING TITLE: 1F FRAMING PLAN\nDRAWING TITLE: 2F FRAMING PLAN\nDRAWING NO: ST-101')
        self.assertEqual(manifest['pdf_pages'][0]['identity_status'], 'conflict')

    def test_explicit_index_title_wins_over_schedule_phrase(self):
        manifest = self.manifest('DRAWING TITLE: DRAWING INDEX / COLUMN SCHEDULE\nDRAWING NO: ST-101')
        self.assertEqual(manifest['pdf_pages'][0]['role'], 'drawing_index')
        self.assertEqual(manifest['context_pages'], [])

    def test_truncated_cad_text_cannot_create_a_unique_link(self):
        cad = {'x.dxf': {'source_member_sha': 'cad',
                 'texts': ['DRAWING TITLE: FRAMING PLAN', 'DRAWING NO: ST-101'],
                 'text_coverage': {'total': 5001, 'retained': 5000, 'truncated': True}}}
        manifest = self.manifest('DRAWING TITLE: FRAMING PLAN\nDRAWING NO: ST-101', cad)
        self.assertEqual(manifest['links'], [])
        self.assertNotEqual(manifest['cad_files'][0]['identity_status'], 'identified')

    def test_malformed_cad_source_path_stays_unresolved(self):
        from core.quantity_sources import attach_member_sources
        for path in [[], {}, 5]:
            with self.subTest(path=path):
                result = attach_member_sources({'columns': [{'source_file': path}]}, self.manifest(''))
                item = result['columns'][0]
                self.assertEqual(item['source_refs'], [])
                self.assertNotEqual(item['source_review']['status'], 'located')

    def test_malformed_category_is_disclosed_without_crashing(self):
        from core.quantity_sources import attach_member_sources
        for category in [5, {}, 'bad']:
            with self.subTest(category=category):
                result = attach_member_sources({'columns': category}, self.manifest(''))
                self.assertEqual(result['columns'], [])
                self.assertTrue(result['notes'])

    def test_invalid_bounding_boxes_are_not_deduplicated(self):
        for box in [
             {'page': 1, 'x': 0, 'y': 0, 'w': 0, 'h': 0},
             {'page': 1, 'x': 0, 'y': 0, 'w': -1, 'h': 1},
             {'page': 1, 'x': .9, 'y': .9, 'w': .5, 'h': .5},
             {'page': 1, 'x': True, 'y': 0, 'w': .1, 'h': .1}]:
            with self.subTest(box=box):
                member = {'mark': 'C1', 'count': 2, 'bbox': box}
                rows, removed, flags = views._dedupe_category_items([deepcopy(member), deepcopy(member)], 'columns')
                self.assertEqual(len(rows), 2)
                self.assertEqual(removed, 0)

    def test_page_only_source_is_not_a_valid_occurrence_location(self):
        from core.quantity_sources import attach_member_sources
        for box in [{'page': 1}, {'page': 1, 'x': 0, 'y': 0, 'w': 0, 'h': 0}]:
            with self.subTest(box=box):
                result = attach_member_sources({'columns': [{'bbox': box}]},
                               self.manifest('DRAWING TITLE: FRAMING PLAN\nDRAWING NO: ST-101'))
                item = result['columns'][0]
                self.assertEqual(item['source_refs'][0]['page'], 1)
                self.assertNotEqual(item['source_review']['status'], 'located')
